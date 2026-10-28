"""Build the new static cable catalog from the old Korda catalog and approved prices."""

from copy import deepcopy
from decimal import Decimal
from html import escape
from pathlib import Path
from urllib.parse import quote, urljoin, urlparse
from urllib.request import Request, urlopen
import hashlib
import re
import zipfile

from docx import Document
from lxml import html
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1] / "dist"
ASSETS = ROOT / "assets" / "cable-catalog"
ASSETS.mkdir(parents=True, exist_ok=True)
BASE = "https://www.korda.su"
SOURCE = "/catalog/grejushhij-kabel/"
PRICES = Path("/Users/ilyaosipov/Downloads/Розница для сайта кабель 27,10,26.xlsx")
FHC_DOC = Path("/Users/ilyaosipov/Downloads/Греющий_кабель_FHC30_страница_на_сайт.docx")
AGENT = {"User-Agent": "Mozilla/5.0"}
cache = {}


def get(url):
    url = urljoin(BASE, url)
    if url not in cache:
        parsed = urlparse(url)
        safe = parsed._replace(path=quote(parsed.path, safe="/"), query=quote(parsed.query, safe="=&%"))
        cache[url] = urlopen(Request(safe.geturl(), headers=AGENT), timeout=30).read()
    return cache[url]


def local_asset(url):
    url = urljoin(BASE, url.strip().removesuffix("%20"))
    suffix = Path(urlparse(url).path).suffix.lower()
    if suffix not in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".pdf", ".doc", ".docx", ".zip"}:
        return url
    name = hashlib.sha1(url.encode()).hexdigest()[:12] + suffix
    dest = ASSETS / name
    if not dest.exists():
        try:
            dest.write_bytes(get(url))
        except Exception as exc:
            print("asset unavailable", url, exc)
            return url
    return "assets/cable-catalog/" + name


def price_text(value):
    n = Decimal(str(value))
    return f"{n:,.2f}".replace(",", " ").replace(".", ",").removesuffix(",00") + " ₽"


def old_price_rows():
    book = load_workbook(PRICES, read_only=True, data_only=True)
    doc = Document(FHC_DOC)
    families = {key: [] for key in ["gws", "gws-cr", "isr-ct", "rgs-cr", "sms-cr", "tms-cr", "tms-ct", "vms-cx", "vms-ct", "hpi13-2ct", "lavita-roll", "lavita-mat"]}
    for row in book["Самрег_250207"].iter_rows(min_row=4, values_only=True):
        model, article, price = row[1], row[9], row[10]
        if not model or not isinstance(price, (int, float)): continue
        key = re.sub(r"\s+", "", str(model).upper())
        if key.startswith("GWS"): family = "gws-cr" if "CR" in key else "gws"
        elif key.startswith("ISR"): family = "isr-ct"
        elif key.startswith("RGS") and "CR" in key: family = "rgs-cr"
        elif key.startswith("SMS"): family = "sms-cr"
        elif key.startswith("TMS"): family = "tms-ct" if "CT" in key else "tms-cr"
        elif key.startswith("VMS"): family = "vms-ct" if "CT" in key else "vms-cx"
        elif key.startswith("HPI"): family = "hpi13-2ct"
        else: continue
        families[family].append((str(model).strip(), str(article), str(row[6] or ""), price, "м"))
    for sheet, family, model_col, article_col, detail_col, price_col in [
        ("Резистив комплект_241011", "lavita-mat", 3, 6, 7, 10),
        ("Резистив_250207", "lavita-roll", 2, 4, 5, 7),
    ]:
        for row in book[sheet].iter_rows(min_row=4, values_only=True):
            model, price = row[model_col], row[price_col]
            if not model or not isinstance(price, (int, float)): continue
            families[family].append((str(model).split()[0], str(row[article_col]), str(row[detail_col] or ""), price, "комплект"))
    for slug, source, model in [
        ("lavita-smart-heat-gws16-2cr", "gws-cr", "GWS 16-2 CR"),
        ("lavita-smart-heat-rgs40-2cr", "rgs-cr", "RGS 40-2 CR"),
        ("lavita-hpi13-2ct", "hpi13-2ct", "HPI13-2CT"),
    ]:
        families[slug] = [row for row in families[source] if row[0] == model]
        assert len(families[slug]) == 1, (slug, model)
    families["lavita-fhc-30"] = [(cells[1], cells[3], cells[2] + " м · " + cells[4] + " Вт", Decimal(cells[5].replace(" ", "").replace(",", ".")), "комплект") for row in doc.tables[1].rows[1:] if (cells := [cell.text.strip() for cell in row.cells]) and cells[5]]
    assert len(families["lavita-fhc-30"]) == 11
    return families, doc


def fhc_content(doc):
    parts = ["<h2>Резистивный греющий кабель Lavita FHC-30</h2>"]
    parts += [f"<p>{escape(p.text.strip())}</p>" for p in doc.paragraphs[2:4] if p.text.strip()]
    with zipfile.ZipFile(FHC_DOC) as z:
        image = z.read("word/media/image1.jpeg")
    image_name = "fhc30-construction.jpeg"
    (ASSETS / image_name).write_bytes(image)
    parts += ["<div class='row'><div class='col-md-6'><h3>Конструкция кабеля Lavita FHC-30</h3>", f"<img src='assets/cable-catalog/{image_name}' alt='Конструкция кабеля Lavita FHC-30'></div><div class='col-md-6'><h3>Технические характеристики</h3><table>"]
    for row in doc.tables[0].rows:
        parts.append("<tr>" + "".join(f"<td>{escape(c.text.strip())}</td>" for c in row.cells) + "</tr>")
    parts += ["</table></div></div>", "<h3>Модели</h3><table><thead><tr><th>Модель</th><th>Длина, м</th><th>Артикул</th><th>Мощность, Вт</th><th>Цена, ₽</th></tr></thead><tbody>"]
    for row in doc.tables[1].rows[1:]:
        c = [x.text.strip() for x in row.cells]
        parts.append("<tr>" + "".join(f"<td>{escape(c[i])}</td>" for i in (1, 2, 3, 4, 5)) + "</tr>")
    parts += ["</tbody></table>", "<h3>Области применения</h3>", "<p>Подогрев ступеней, лестниц и входов в здания; обогрев открытых площадок, кровли и водостоков. Укладка в бетонную стяжку для обычных сред.</p>"]
    return "".join(parts)


def sanitized_content(tree, slug, doc):
    if slug == "lavita-fhc-30": return fhc_content(doc)
    roots = tree.xpath("//h2")
    if not roots: raise ValueError(f"Missing product description: {slug}")
    root = deepcopy(roots[0].getparent())
    for bad in root.xpath(".//form|.//script|.//iframe|.//input|.//button|.//select|.//noscript"):
        bad.drop_tree()
    for bad in root.xpath(".//h1|.//h3"):
        if bad.tag == "h1" or bad.text_content().strip().startswith(("Цена кабеля", "Цена плёнки")):
            bad.drop_tree()
    for stale_price in root.xpath(".//p[starts-with(normalize-space(.), 'Стоимость комплекта =')]"):
        stale_price.drop_tree()
    for el in root.iter():
        if not isinstance(el.tag, str): continue
        for attr in list(el.attrib):
            if attr not in {"href", "src", "alt", "class"}: del el.attrib[attr]
        if el.tag == "img":
            el.set("src", local_asset(el.get("src", "")))
            el.set("loading", "lazy")
        elif el.tag == "a" and el.get("href"):
            href = el.get("href").strip().removesuffix("%20")
            if href.startswith("#"): continue
            if href.lower().split("?")[0].endswith((".pdf", ".doc", ".docx", ".zip", ".jpg", ".jpeg", ".png")):
                el.set("href", local_asset(href))
            elif "/catalog/grejushhij-kabel/" in href:
                el.set("href", "cable-" + Path(urlparse(href).path).stem + ".html")
            else:
                el.set("href", urljoin(BASE, href))
    return "".join(html.tostring(child, encoding="unicode", method="html") for child in root)


def table_html(rows):
    if not rows: return "<p class='price-unavailable'>Цена по запросу</p>"
    head = "<table class='cable-price-table'><thead><tr><th>Модель</th><th>Артикул</th><th>Параметр</th><th>Цена</th></tr></thead><tbody>"
    body = "".join(f"<tr><td>{escape(model)}</td><td>{escape(article)}</td><td>{escape(detail)}</td><td>{price_text(price)} / {unit}</td></tr>" for model, article, detail, price, unit in rows)
    return head + body + "</tbody></table>"


def main():
    prices, doc = old_price_rows()
    source = html.fromstring(get(SOURCE))
    labels = {}
    for el in source.xpath("//input[@data-filter]"):
        label = source.xpath("//label[@for=$id]", id=el.get("id"))
        if label:
            labels[el.get("data-filter").lstrip(".")] = label[0].text_content().strip()
    template = html.fromstring((ROOT / "household-solutions.html").read_text())
    header = html.tostring(template.xpath("//header[@class='site-header']")[0], encoding="unicode", method="html")
    footer = html.tostring(template.xpath("//footer[@class='site-footer']")[0], encoding="unicode", method="html")
    header = header.replace('href="index.html#product"', 'href="heating-cable.html"').replace('href="household-solutions.html"', 'href="household-solutions.html"')
    header = header.replace('class="nav-trigger active-nav" href="household-solutions.html"', 'class="nav-trigger" href="household-solutions.html"')
    header = header.replace('class="active-nav" href="household-solutions.html"', 'href="household-solutions.html"')
    header = header.replace('href="heating-cable.html">Подобрать кабель', 'class="active-nav" href="heating-cable.html">Подобрать кабель')
    footer = footer.replace('href="index.html#product"', 'href="heating-cable.html"')
    frame = """<!doctype html><html lang='ru'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><meta name='robots' content='index,follow'><title>{title} | КОРДА</title><meta name='description' content='{description}'><link rel='preconnect' href='https://fonts.googleapis.com'><link rel='preconnect' href='https://fonts.gstatic.com' crossorigin><link href='https://fonts.googleapis.com/css2?family=Montserrat:wght@400;500;600;700;800&display=swap' rel='stylesheet'><link rel='stylesheet' href='assets/styles.css?v=20261027-buttons'><link rel='stylesheet' href='assets/industrial-pages.css?v=20261027-buttons'><link rel='stylesheet' href='assets/heating-cable.css?v=20261028-cards'></head><body class='inner-page'><a class='skip-link' href='#main'>К основному содержанию</a>{header}<main id='main' class='inner-main'>{body}</main>{footer}<script src='assets/app.js' defer></script>{script}</body></html>"""
    items = []
    for li in source.xpath("//ul[@id='cabeli']/li"):
        heading = li.xpath(".//div[contains(@class,'hidden-xs')]//h1")
        link = li.xpath(".//a[contains(@class,'btn')]")
        thumb = li.xpath(".//img")
        if not heading or not link or not thumb: continue
        name = heading[0].text_content().strip()
        old_url = urljoin(BASE, link[0].get("href"))
        slug = Path(urlparse(old_url).path).stem
        klasses = li.get("class", "").split()
        category = next((x for x in klasses if x in {"ct1", "ct2", "ct3"}), "ct1")
        filters = [x for x in klasses if x in labels]
        manufacturer = next((x for x in klasses if x in {"Lavita", "Ace_Floor", "ТЕРМ"}), "")
        intro = li.xpath(".//p")
        description = intro[0].text_content().split("Производитель:")[0].strip() if intro else name
        image = local_asset(thumb[0].get("src"))
        old_detail = html.fromstring(get(old_url))
        body = sanitized_content(old_detail, slug, doc)
        rows = prices.get(slug, [])
        shown_price = "от " + price_text(min(x[3] for x in rows)) + " / " + rows[0][4] if rows else "Цена по запросу"
        price_note = "Цены из переданного описания FHC-30. Перед заказом уточняйте актуальность." if slug == "lavita-fhc-30" else ("Цены из прайса от 27.10.2026. Для остальных моделей стоимость по запросу." if rows else "Позиции нет в прайсе от 27.10.2026; цену уточняйте перед заказом.")
        detail_main = f"<div class='cable-crumb shell'><a href='heating-cable.html'>Подобрать кабель</a> / {escape(name)}</div><section class='cable-detail-hero shell'><div><p class='eyebrow'>Греющий кабель · {escape(manufacturer.replace('_',' '))}</p><h1>{escape(name)}</h1><p>{escape(description)}</p><strong>{shown_price}</strong><div><a class='button button-primary' href='mailto:zayavka@korda.spb.ru?subject={quote('Запрос по кабелю '+name)}'>Запросить подбор →</a></div></div><img src='{image}' alt='{escape(name)}'></section><div class='cable-detail shell'><section class='cable-specs'><h2>Модели и цены</h2>{table_html(rows)}<p class='cable-source-note'>{price_note}</p></section><section class='cable-description'>{body}<p class='cable-source-note'>Документы перенесены из прежнего каталога. Перед использованием проверяйте срок их действия.</p></section></div>"
        (ROOT / f"cable-{slug}.html").write_text(re.sub(r"[ \t]+(?=\n)", "", frame.format(title=escape(name), description=escape(description[:150], quote=True), header=header, body=detail_main, footer=footer, script="")), encoding="utf-8")
        items.append(dict(name=name, slug=slug, category=category, filters=filters, manufacturer=manufacturer, description=description, image=image, price=shown_price))
        print("product", name, len(rows), len(body))
    filter_groups = [("Область применения", "app"), ("Условия применения", "env"), ("Производитель", "maker")]
    filters_html = ""
    for title, prefix in filter_groups:
        options = [(key, label) for key, label in labels.items() if key.startswith(prefix)] if prefix != "maker" else [(x, x.replace('_', ' ')) for x in ["Lavita", "Ace_Floor", "ТЕРМ"]]
        filters_html += f"<fieldset><legend>{title}</legend>" + "".join(f"<label><input type='checkbox' name='filter' data-group='{prefix}' value='{escape(key)}'><span>{escape(label)}</span></label>" for key, label in options) + "</fieldset>"
    cards = "".join(f"<article class='cable-card' data-category='{x['category']}' data-filters='{' '.join(x['filters'] + ([x['manufacturer']] if x['manufacturer'] else []))}'><a class='cable-card__image' href='cable-{x['slug']}.html'><img src='{x['image']}' alt='{escape(x['name'])}' loading='lazy'></a><div><span class='cable-card__maker'>{escape(x['manufacturer'].replace('_',' '))}</span><h2>{escape(x['name'])}</h2><p>{escape(x['description'])}</p><strong>{x['price']}</strong><a class='cable-card__link' href='cable-{x['slug']}.html'>Подробнее →</a></div></article>" for x in items)
    catalog_body = f"<section class='cable-hero shell'><p class='eyebrow'>Каталог КОРДА</p><h1>Подобрать греющий кабель</h1><p>Выберите тип кабеля и условия применения.</p></section><div class='cable-layout shell'><aside class='cable-filters'><h2>Фильтры</h2><button type='button' id='filter-toggle' aria-expanded='false' aria-controls='cable-filter-options'>Показать фильтры</button><button type='button' id='clear-filters'>Сбросить</button><div id='cable-filter-options'>{filters_html}</div></aside><section class='cable-results'><div class='cable-tabs' role='tablist' aria-label='Тип кабеля'><button type='button' class='active' data-category='all' aria-selected='true'>Все</button><button type='button' data-category='ct1' aria-selected='false'>Саморегулирующийся кабель</button><button type='button' data-category='ct2' aria-selected='false'>Резистивный кабель</button><button type='button' data-category='ct3' aria-selected='false'>Инфракрасная плёнка</button></div><p id='cable-count' aria-live='polite'></p><div class='cable-grid'>{cards}</div><p class='cable-empty' hidden>По выбранным фильтрам товаров нет.</p></section></div>"
    (ROOT / "heating-cable.html").write_text(frame.format(title="Греющий кабель", description="Каталог греющего кабеля КОРДА: саморегулирующийся, резистивный и инфракрасная плёнка.", header=header, body=catalog_body, footer=footer, script="<script src='assets/heating-cable.js' defer></script>"), encoding="utf-8")
    assert len(items) == len(source.xpath("//ul[@id='cabeli']/li")) == 20
    assert len(list(ROOT.glob("cable-*.html"))) >= 20
    print("built", len(items), "products")


if __name__ == "__main__": main()
