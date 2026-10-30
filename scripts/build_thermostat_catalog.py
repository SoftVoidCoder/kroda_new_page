"""Import all thermostat cards and technical content; use approved spreadsheet prices."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from html import escape as e
from pathlib import Path
from urllib.parse import urljoin, urlparse, quote
from lxml import html
from openpyxl import load_workbook
from docx import Document
from zipfile import ZipFile
import hashlib
import re
import sys
sys.path.insert(0, str(Path(__file__).parent))
from build_cable_catalog import get, price_text, ROOT, BASE

ASSETS = ROOT / 'assets/thermostats'
ASSETS.mkdir(exist_ok=True)
PRICE_FILE = Path('/Users/ilyaosipov/Downloads/Розница_для_сайта_темостаты_29,10,26 (1).xlsx')
KIT_I_DOC = Path('/Users/ilyaosipov/Downloads/Термостат KIT-I( на сайт).docx')
KIT_I_PDF = Path('/Users/ilyaosipov/Downloads/진_Thermostat_Kit_설명서_ТЕРМОРЕГУЛЯТОР_THERMOSTAT_KIT_rev_260923.pdf')
SOURCE = BASE + '/termostaty/'
SNAPSHOT = Path(__file__).parent / 'thermostat-source'
SNAPSHOT.mkdir(exist_ok=True)

def fetch(url):
    path = SNAPSHOT / (hashlib.sha1(url.encode()).hexdigest()[:12] + '.html')
    if not path.exists(): path.write_bytes(get(url))
    return html.fromstring(path.read_bytes())

def asset(url, base):
    absolute = urljoin(base, url.strip())
    suffix = Path(urlparse(absolute).path).suffix.lower()
    if suffix not in {'.jpg','.jpeg','.png','.gif','.webp','.svg','.pdf','.doc','.docx','.zip','.xls','.xlsx'}: return absolute
    dest = ASSETS / (hashlib.sha1(absolute.encode()).hexdigest()[:12] + suffix)
    if not dest.exists(): dest.write_bytes(get(absolute))
    return 'assets/thermostats/' + dest.name

def page_name(url): return 'thermostat-' + Path(urlparse(url).path).stem.lower() + '.html'

def content(tree, url, urls, original_images):
    roots = tree.xpath('//div[contains(concat(" ",normalize-space(@class)," ")," tpl-block-full ")]')
    assert len(roots) == 1, url
    root = deepcopy(roots[0])
    for bad in root.xpath('.//form|.//script|.//iframe|.//input|.//button|.//select|.//noscript|.//comment()'):
        bad.getparent().remove(bad)
    # Product name and photo already appear in the page header.
    first_heading = root.xpath('.//h2')[0]
    first_heading.getparent().getparent().drop_tree()
    first_photo = root.xpath('.//img')[0]
    if 'col-md-3' in first_photo.getparent().get('class', ''):
        first_photo.getparent().drop_tree()
    for p in root.xpath('.//p[contains(@class,"lead")]'):
        if p.text_content().strip().startswith('Цена:'): p.drop_tree()
    for col in root.xpath('.//div[contains(@class,"col-md-3")]'):
        if not col.text_content().strip() and not col.xpath(".//img"): col.drop_tree()
    for node in root.iter():
        for attr in ('text', 'tail'):
            text = getattr(node, attr)
            if text: setattr(node, attr, re.sub(r'\d[\d \u00a0]*(?:[,.]\d+)?\s*руб\.?', 'цена по запросу', text))
    for el in root.iter():
        if not isinstance(el.tag,str): continue
        for attr in list(el.attrib):
            if attr not in {'href','src','alt','class','colspan','rowspan','id'}: del el.attrib[attr]
        if el.tag == 'img':
            src = urljoin(url, el.get('src'))
            el.set('src', asset(original_images.get(src, src), url)); el.set('loading','lazy')
            if not el.get('alt'): el.set('alt','Техническая иллюстрация')
        if el.tag == 'a' and el.get('href'):
            href = el.get('href').strip()
            absolute = urljoin(url,href)
            if href.startswith('#'): continue
            if absolute in urls: el.set('href',page_name(absolute))
            else: el.set('href',asset(absolute,url))
    assert len(root.xpath('.//table')) == len(roots[0].xpath('.//table')), url
    assert len(root.xpath('.//img')) + 1 == len(roots[0].xpath('.//img')), url
    assert len(root.xpath('.//a[contains(@href,".pdf")]')) == len(roots[0].xpath('.//a[contains(@href,".pdf")]')), url
    # Old Bootstrap columns are laid out by the existing technical-page stylesheet.
    return re.sub(r'[ \t]+(?=\n)', '', html.tostring(root,encoding='unicode'))

def prices():
    rows=list(load_workbook(PRICE_FILE,read_only=True,data_only=True).active.values)
    mapping={'RTC70.26':'rtc70-26','E51.716':'e51-716','E91.716':'e_91_716','UTH-170(4KW)':'uth-170','UTH-150(2KW) B type':'uth150','UTH-150(2KW) A type':'uth150-2','UTH-120(4KW)':'uth120','Термостат Кит WP влагозащищенный(черный)':'termostat-kit-lavita','Термостат KIT -I (без влагозащиты, желтый)':'kit-i'}
    result={}
    for row in rows:
        for value in row:
            if str(value).strip() in mapping:
                numbers=[v for v in row if isinstance(v,(int,float))]
                # The retail price is the last numeric cell in each approved row.
                result[mapping[str(value).strip()]]=numbers[-1]
    assert len(result)==9,result
    assert result['termostat-kit-lavita']==3700 and result['kit-i']==1700,result
    assert result['rtc70-26']==700 and result['uth150']==2410,result
    return result


def kit_i_content():
    doc = Document(KIT_I_DOC)
    with ZipFile(KIT_I_DOC) as source:
        (ASSETS / 'kit-i-original.jpeg').write_bytes(source.read('word/media/image1.jpeg'))
    (ASSETS / 'kit-i-passport.pdf').write_bytes(KIT_I_PDF.read_bytes())
    headings = {'Как это работает', 'Комплектация', 'Особенности монтажа'}
    parts = []
    table_open = False
    for paragraph in doc.paragraphs[1:]:
        text = paragraph.text.strip()
        if not text or text == '(Паспорт отдельным файлом)': continue
        if text.startswith('Технические'):
            parts.append('<h2>Технические характеристики</h2><table><tbody>')
            table_open = True
        elif text in headings:
            if table_open:
                parts.append('</tbody></table>')
                table_open = False
            parts.append('<h2>' + e(text) + '</h2>')
        elif table_open:
            label, value = text.split(':', 1)
            parts.append('<tr><th>' + e(label) + '</th><td>' + e(value.strip()) + '</td></tr>')
        else:
            parts.append('<p>' + e(text) + '</p>')
    assert not table_open
    parts.append("<h2>Документация</h2><p><a href='assets/thermostats/kit-i-passport.pdf'>Паспорт и инструкция по эксплуатации THERMOSTAT KIT (PDF)</a></p>")
    assert sum(p.text.strip().startswith('Напряжение:') for p in doc.paragraphs) == 1
    return ''.join(parts)

def main():
    source=fetch(SOURCE)
    blocks=source.xpath('//div[@id="filter"]/div[contains(@class,"sub-block")]')
    assert len(blocks)==39
    urls=[urljoin(SOURCE,b.xpath('.//a')[0].get('href')) for b in blocks]
    with ThreadPoolExecutor(max_workers=6) as pool: details=list(pool.map(fetch,urls))
    original_images = {urljoin(url, block.xpath('.//img')[0].get('src')): urljoin(url, tree.xpath('//div[contains(@class,"tpl-block-full")]//img')[0].get('src')) for block, url, tree in zip(blocks, urls, details)}
    labels={a.get('data-filter'):a.text_content().strip().replace('Tеплый','Тёплый') for a in source.xpath('//a[@data-filter]') if a.get('data-filter')!='all'}
    makers={'607','599','598','597','319','315'}
    approved=prices()
    frame_source=(ROOT/'heating-cable.html').read_text()
    frame_tree=html.fromstring(frame_source)
    header=html.tostring(frame_tree.xpath('//header')[0],encoding='unicode')
    header=header.replace('class="active-nav" href="heating-cable.html"','href="heating-cable.html"')
    header=re.sub(r'href="[^"]*">Термостаты', 'class="active-nav" href="thermostats.html">Термостаты',header)
    footer=html.tostring(frame_tree.xpath('//footer')[0],encoding='unicode')
    head=frame_source[:frame_source.index('<body')]
    head=re.sub(r'<title>.*?</title>','<title>{title} | КОРДА</title>',head)
    head=re.sub(r"<meta name=['\"]description['\"] content=['\"][^'\"]*['\"]>",'<meta name="description" content="{description}">',head)
    head=head.replace('</head>',"<link rel='stylesheet' href='assets/thermostats.css?v=20261029-originals'></head>")
    def render(title,desc,body,js=''):
        return head.replace('{title}',e(title)).replace('{description}',e(desc,quote=True))+"<body class='inner-page thermostats-page'><a class='skip-link' href='#main'>К основному содержанию</a>"+header+"<main id='main' class='inner-main'>"+body+'</main>'+footer+"<script src='assets/app.js' defer></script>"+js+'</body></html>'
    items=[]
    for block,url,tree in zip(blocks,urls,details):
        name=block.xpath('.//h1')[0].text_content().strip()
        slug=Path(urlparse(url).path).stem.lower()
        tags=[c for c in block.get('class').split() if c in labels]
        maker=next(labels[c] for c in tags if c in makers)
        image=asset(original_images[urljoin(url, block.xpath('.//img')[0].get('src'))],url)
        shown=price_text(approved[slug]) if slug in approved else 'Цена по запросу'
        body=content(tree,url,set(urls),original_images)
        full_image=asset(tree.xpath('//div[contains(@class,"tpl-block-full")]//img')[0].get('src'),url)
        summary = ''
        if slug == 'termostat-kit-lavita':
            name = 'Термостат KIT WP — влагозащищённый'
            summary = '<p>Влагозащищённый термостат Lavita. Автоматическое включение при температуре ниже +5 °C, выключение — выше +15 °C. Максимальная мощность — 2 кВт.</p>'
            body = body.replace('Диапазон регулирования температуры', 'Автоматическое выключение при температуре окружающей среды')
        detail=f"<div class='cable-crumb shell'><a href='thermostats.html'>Термостаты</a> / {e(name)}</div><section class='cable-detail-hero shell'><div><p class='eyebrow'>{e(maker)}</p><h1>{e(name)}</h1>{summary}<strong>{shown}</strong><a class='button button-primary' href='mailto:zayavka@korda.spb.ru?subject={quote('Запрос: '+name)}'>Запросить подбор →</a></div><img src='{full_image}' alt='{e(name)}'></section><div class='cable-detail shell'><section class='cable-description'>{body}</section></div>"
        (ROOT/page_name(url)).write_text(render(name,'Характеристики, описание и документация: '+name,detail))
        items.append(dict(name=name,url=page_name(url),image=image,tags=tags,maker=maker,price=shown,summary=('Влагозащищённый. Включение ниже +5 °C, выключение выше +15 °C.' if slug == 'termostat-kit-lavita' else '')))
        if not items[-1]['summary']: del items[-1]['summary']
        print('imported',slug,shown,flush=True)
    kit_name = 'Термостат KIT-I — без влагозащиты'
    kit_image = 'assets/thermostats/kit-i-original.jpeg'
    kit_price = price_text(approved['kit-i'])
    kit_body = kit_i_content()
    kit_summary = 'Для установки в сухом месте с защитой от воды и конденсата. Включение ниже +5 °C, выключение при достижении +15 °C. Максимальная мощность — 2 кВт.'
    kit_detail = f"<div class='cable-crumb shell'><a href='thermostats.html'>Термостаты</a> / {e(kit_name)}</div><section class='cable-detail-hero shell'><div><p class='eyebrow'>Lavita</p><h1>{e(kit_name)}</h1><p>{kit_summary}</p><strong>{kit_price}</strong><a class='button button-primary' href='mailto:zayavka@korda.spb.ru?subject={quote('Запрос: KIT-I')}'>Запросить подбор →</a></div><img src='{kit_image}' alt='Жёлтый термостат KIT-I Lavita без влагозащиты'></section><div class='cable-detail shell'><section class='cable-description'>{kit_body}</section></div>"
    (ROOT/'thermostat-kit-i.html').write_text(render(kit_name,kit_summary,kit_detail))
    items.insert(15,dict(name=kit_name,url='thermostat-kit-i.html',image=kit_image,tags=['315','325','327','397'],maker='Lavita',price=kit_price,summary='Без влагозащиты. Включение ниже +5 °C, выключение при достижении +15 °C.'))
    mock=html.fromstring((Path(__file__).resolve().parents[1] / 'docs/termostaty-mockup.html').read_text())
    intro=mock.xpath('//section')[0]
    intro.set('class','thermo-section shell')
    # The approved mockup supplies layout and copy; real product photos illustrate the applications.
    for card,item_index in zip(intro.xpath('.//div[@class="selection"]|.//article'),[0,5,1,8,9]):
        picture=html.Element('img',src=items[item_index]['image'],alt=items[item_index]['name'],loading='lazy')
        card.append(picture)
    guides=html.tostring(intro,encoding='unicode')
    filter_html=''
    for title,group,keys in [('Производитель','maker',makers),('Область применения','app',set(labels)-makers)]:
        filter_html+=f'<fieldset><legend>{title}</legend>'
        for key,label in labels.items():
            if key in keys: filter_html+=f"<label><input type='checkbox' data-group='{group}' value='{key}'><span>{e(label)}</span></label>"
        filter_html+='</fieldset>'
    cards=''.join(f"<article class='cable-card' data-filters='{' '.join(x['tags'])}'><a class='cable-card__image' href='{x['url']}'><img src='{x['image']}' alt='{e(x['name'])}' loading='lazy'></a><div><span class='cable-card__maker'>{e(x['maker'])}</span><h3><a href='{x['url']}'>{e(x['name'])}</a></h3>{('<p>'+e(x['summary'])+'</p>') if x.get('summary') else ''}<strong>{x['price']}</strong><a class='cable-card__link' href='{x['url']}'>Подробнее →</a></div></article>" for x in items)
    catalog=f"<div class='cable-layout shell' id='catalog'><aside class='cable-filters'><h2>Фильтры</h2><button type='button' id='filter-toggle' aria-expanded='false' aria-controls='cable-filter-options'>Показать фильтры</button><button type='button' id='clear-filters'>Сбросить</button><div id='cable-filter-options'>{filter_html}</div></aside><section class='cable-results' aria-label='Каталог термостатов'><h2>Термостаты и датчики</h2><p id='cable-count' aria-live='polite'>Найдено: 40</p><div class='cable-grid'>{cards}</div><p class='cable-empty' hidden>По выбранным фильтрам товаров нет.</p></section></div>"
    (ROOT/'thermostats.html').write_text(render('Электронные термостаты для греющего кабеля','Термостаты и датчики Lavita, OJ Electronics, Uriel Electronics, ТЕРМ, ПЭЛЗ и АЛТИМЕР. Подбор по производителю и области применения.',guides+catalog,"<script src='assets/heating-cable.js' defer></script>").replace('thermostats.css?v=20261029-originals', 'thermostats.css?v=20261030-alignment'))
    # Link the section from every existing shared navigation and footer.
    for path in ROOT.glob('*.html'):
        txt=path.read_text()
        changed=re.sub(r'href=([\"\'])[^\"\']*\1(?=>Термостаты)', 'href="thermostats.html"',txt)
        if changed!=txt: path.write_text(changed)
    import json
    (ASSETS/'products.json').write_text(json.dumps(items,ensure_ascii=False,indent=2))
    assert len({x['url'] for x in items})==40
    assert sum(x['price']!='Цена по запросу' for x in items)==9
    print('Built 40 products, 9 approved prices.')

if __name__=='__main__': main()
